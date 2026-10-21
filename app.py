"""Data QA Agent: ask questions about Excel and CSV files."""
import asyncio
import inspect
import json
import os
import time

import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

from agent.context_budget import DEFAULT_INPUT_TOKENS
from agent.factory import DEFAULT_MODEL, GROQ_KEY_URL, MODELS, build_model
from agent.qa_agent import answer_question
from services.analysis import friendly_error
from services.context import bounded_chat
from services.data_tools import MAX_TOTAL_BYTES, SUPPORTED_TYPES, excel_previews, read_upload, sanitize_dataframe
from services.workbook import Workbook, column_kind


def _width_kwarg(stretch: bool = True) -> dict:
    if 'width' in inspect.signature(st.button).parameters:
        return {'width': 'stretch' if stretch else 'content'}
    return {'use_container_width': stretch}


WIDE = _width_kwarg(True)
SECRET_NAME = 'GROQ_API_KEY'


def load_key():
    try:
        secret = st.secrets.get(SECRET_NAME, '')
    except (FileNotFoundError, st.errors.StreamlitSecretNotFoundError):
        secret = ''
    if secret or os.environ.get(SECRET_NAME):
        return secret or os.environ[SECRET_NAME]
    if os.environ.get('ENABLE_OS_KEYRING') != '1':
        return ''
    try:
        import keyring
        return keyring.get_password('PersonalDataAgent', 'groqcloud') or ''
    except Exception:
        return ''


def close_workbook():
    if st.session_state.get('workbook'):
        st.session_state.workbook.close()
    st.session_state.workbook = None
    st.session_state.dataset_signature = None
    st.session_state.messages = []


def new_chat():
    st.session_state.messages = []


def display(frame):
    """Arrow-safe copy for display; date-only columns drop the 00:00:00 time."""
    frame = frame.copy()
    for column in frame.columns:
        values = frame[column]
        if pd.api.types.is_datetime64_any_dtype(values) and (values.dropna() == values.dropna().dt.normalize()).all():
            frame[column] = values.dt.date
    return sanitize_dataframe(frame)


def ask(question):
    st.session_state.pending_question = question


def suggestions(workbook):
    """Generic starter questions built from column types, never from field-name lists."""
    sid, frame = next(iter(workbook.frames.items()))
    kinds = {column: column_kind(frame[column]) for column in frame.columns[:40]}
    # Measures vary but repeat; all-unique numbers are usually identifiers.
    numbers = [c for c, k in kinds.items() if k == 'number' and 10 < frame[c].nunique() < len(frame)]
    categories = [c for c, k in kinds.items() if k == 'text' and 1 < frame[c].nunique() <= 30]
    dates = [c for c, k in kinds.items() if k == 'date']
    questions = [f'Give me a quick overview of {workbook.labels[sid]}', 'Which columns have missing values?']
    if numbers and categories:
        questions.append(f'What is the total {numbers[0]} by {categories[0]}?')
    if numbers and dates:
        questions.append(f'How does {numbers[0]} change by month?')
    elif numbers:
        questions.append(f'Show the top 5 rows by {numbers[0]}')
    return questions[:4]


def render_steps(steps, key, expanded=False):
    if not steps:
        return
    tables = [s for s in steps if s.frame is not None and not s.frame.empty]
    label = f'📊 Data behind this answer · {len(tables)} table{"s" * (len(tables) != 1)}' if tables else '🔎 How this was answered'
    with st.expander(label, expanded=expanded):
        for index, step in enumerate(steps):
            st.markdown(f'**{step.title}**  \n<span style="opacity:.7">{step.summary}</span>', unsafe_allow_html=True)
            if step.sql:
                st.code(step.sql, language='sql')
            if step.frame is None or step.frame.empty:
                continue
            frame = step.frame
            numbers = [c for c in frame.columns[1:] if pd.api.types.is_numeric_dtype(frame[c])]
            if step.chart and numbers:
                table_tab, chart_tab = st.tabs(['Table', 'Chart'])
                chart_data = frame.set_index(frame.columns[0])[numbers[:3]]
                chart_data.index = chart_data.index.astype(str)
                with chart_tab:
                    (st.line_chart if step.chart == 'line' else st.bar_chart)(chart_data)
            else:
                table_tab = st.container()
            with table_tab:
                st.dataframe(display(frame), hide_index=True, **WIDE)
                st.download_button('⬇️ Download CSV', frame.to_csv(index=False).encode('utf-8-sig'),
                                   f'result_{key}_{index + 1}.csv', mime='text/csv', key=f'download_{key}_{index}')


st.set_page_config(page_title='Data QA Agent', page_icon='📊', layout='wide')
st.markdown("""<style>
[data-testid="stMainBlockContainer"], .block-container {max-width: 1100px; padding-top: 2.2rem;}
[data-testid="stSidebar"] h3 {font-size: 1rem; margin-bottom: .2rem;}
div[data-testid="stExpander"] details {border-radius: .6rem;}
</style>""", unsafe_allow_html=True)

for name, default in (('messages', []), ('workbook', None), ('dataset_signature', None), ('pending_question', None)):
    st.session_state.setdefault(name, default)

with st.sidebar:
    st.markdown('### 🤖 Model · GroqCloud')
    model_name = st.selectbox('Model', list(MODELS), key='model', accept_new_options=True,
                              help='Pick a model, or type any model ID enabled on your GroqCloud account.') or DEFAULT_MODEL
    st.caption(MODELS.get(model_name, 'Custom model: it must support tool calling'))
    stored_key = load_key()
    key_input = st.text_input('API key', type='password', key='api_key',
                              placeholder='Using the configured key' if stored_key else 'Paste your GroqCloud key',
                              help='Kept only in this browser session. Never paste it into the chat.')
    api_key = key_input.strip() or stored_key
    st.caption('🟢 Ready' if api_key else f'🔑 [Get a free GroqCloud key]({GROQ_KEY_URL})')
    with st.expander('Advanced'):
        input_budget = st.select_slider('Context budget (tokens per request)', [6000, 12000, 24000, 48000],
                                        value=DEFAULT_INPUT_TOKENS, help='Lower it if you hit GroqCloud rate limits.')
        if os.environ.get('ENABLE_OS_KEYRING') == '1' and key_input.strip() and st.button('Save key to this computer'):
            try:
                import keyring
                keyring.set_password('PersonalDataAgent', 'groqcloud', key_input.strip())
                st.success('Saved in the OS credential store.')
            except Exception as exc:
                st.warning(f'Could not save the key: {exc}')

    st.markdown('### 📁 Data')
    uploads = st.file_uploader('Excel or CSV files', type=SUPPORTED_TYPES, accept_multiple_files=True,
                               help='Up to 20 MB per file and 60 MB in total. Each Excel sheet becomes its own table.')

excel_options = {}
if uploads:
    preview_cache = st.session_state.setdefault('excel_header_previews', {})
    active = {(f.name, f.size, f.file_id) for f in uploads if f.name.lower().endswith(('.xlsx', '.xlsm', '.xls'))}
    for cached in [c for c in preview_cache if c not in active]:
        del preview_cache[cached]
    for upload in uploads:
        if not upload.name.lower().endswith(('.xlsx', '.xlsm', '.xls')):
            continue
        cache_key = (upload.name, upload.size, upload.file_id)
        try:
            if cache_key not in preview_cache:
                preview_cache[cache_key] = excel_previews(upload)
            with st.sidebar.expander(f'Excel header settings: {upload.name}'):
                st.caption('Change the header row if the sheet starts with titles or merged cells. Row numbers match Excel.')
                for sheet, preview in preview_cache[cache_key].items():
                    st.markdown(f'**{sheet}**')
                    key = f'excel_{upload.file_id}_{sheet}'
                    mode = st.selectbox('Header mode', ['Automatic', 'Select header row', 'No header'], key=key + '_mode')
                    row, depth = 'auto', 1
                    if mode == 'Select header row':
                        row = int(st.number_input('Header starts at Excel row', min_value=1, max_value=1000,
                                                  value=1, key=key + '_row')) - 1
                        depth = int(st.number_input('Number of header rows', min_value=1, max_value=3,
                                                    value=1, key=key + '_depth'))
                    elif mode == 'No header':
                        row = None
                    excel_options.setdefault(upload.file_id, {})[sheet] = {'header_row': row, 'header_rows': depth}
                    visible = preview.copy()
                    visible.columns = [get_column_letter(i + 1) for i in range(len(visible.columns))]
                    visible.index = range(1, len(visible) + 1)
                    visible.index.name = 'Excel row'
                    st.dataframe(sanitize_dataframe(visible), **WIDE)
        except Exception as exc:
            st.sidebar.warning(f'Excel header preview unavailable for {upload.name}: {exc}')
    signature = (tuple((f.name, f.size, f.file_id) for f in uploads), json.dumps(excel_options, sort_keys=True))
    if signature != st.session_state.dataset_signature:
        try:
            if sum(f.size for f in uploads) > MAX_TOTAL_BYTES:
                raise ValueError('Combined uploads must be 60 MB or smaller.')
            tables = {}
            with st.spinner('Reading your files…'):
                for upload in uploads:
                    for name, frame in read_upload(upload, excel_headers=excel_options.get(upload.file_id)).items():
                        if name in tables:
                            raise ValueError(f'Duplicate table name: {name}; rename one file.')
                        tables[name] = frame
                        if sum(int(df.memory_usage(deep=True).sum()) for df in tables.values()) > 200 * 1024 * 1024:
                            raise ValueError('The files use more than 200 MB of memory; upload fewer files.')
            close_workbook()
            st.session_state.workbook = Workbook(tables)
            st.session_state.dataset_signature = signature
        except Exception as exc:
            close_workbook()
            st.error(f'Could not load the files: {exc}')
elif st.session_state.workbook:
    close_workbook()
    st.session_state.excel_header_previews = {}

workbook = st.session_state.workbook
scope = None
with st.sidebar:
    if workbook and len(workbook.frames) > 1:
        choice = st.selectbox('Ask about', [None] + list(workbook.frames),
                              format_func=lambda sid: 'All sheets' if sid is None else workbook.labels[sid],
                              help='Focusing on one sheet keeps requests small and answers faster.')
        scope = [choice] if choice else None
    if st.session_state.messages:
        st.divider()
        left, right = st.columns(2)
        left.button('🧹 New chat', on_click=new_chat, **WIDE)
        right.download_button('⬇️ Export', json.dumps([{'role': m['role'], 'content': m['content']}
                                                       for m in st.session_state.messages], ensure_ascii=False, indent=2),
                              'data_qa_chat.json', mime='application/json', **WIDE)

st.title('📊 Data QA Agent')
st.caption('Ask questions about your Excel and CSV files in plain language. The agent runs queries over every row '
           'on this machine and sends only small, relevant results to the AI model.')

if not workbook:
    columns = st.columns(3)
    for column, (title, text) in zip(columns, [
            ('1 · Add your key', 'Paste a GroqCloud API key in the sidebar.'),
            ('2 · Upload data', 'Drop Excel or CSV files in the sidebar. Large sheets are fine.'),
            ('3 · Ask', 'Totals, filters, rankings, trends, missing values, comparisons…')]):
        with column.container(border=True):
            st.markdown(f'**{title}**')
            st.caption(text)
else:
    total_rows = sum(len(frame) for frame in workbook.frames.values())
    with st.expander(f'📄 {len(workbook.frames)} sheet{"s" * (len(workbook.frames) != 1)} loaded · {total_rows:,} rows',
                     expanded=not st.session_state.messages):
        ids = list(workbook.frames)
        tabs = st.tabs([workbook.labels[sid] for sid in ids[:10]]) if len(ids) > 1 else [st.container()]
        for sid, tab in zip(ids, tabs):
            frame = workbook.frames[sid]
            with tab:
                st.caption(f'{len(frame):,} rows × {len(frame.columns)} columns · showing the first 100 rows')
                if frame.attrs.get('notes'):
                    st.caption(f'{len(frame.attrs["notes"])} line(s) of text below the table (notes, signatures) '
                               'are not counted as records.')
                if 'header_row' in frame.attrs:
                    header = frame.attrs['header_row']
                    if frame.attrs.get('generic_headers'):
                        st.caption('Unlabelled columns: ' + ', '.join(frame.attrs['generic_headers'])
                                   + '. If the real labels are on another row, change it in Excel header settings.')
                    elif header and header > 1:
                        st.caption(f'Headers read from Excel row {header}.')
                st.dataframe(display(frame.head(100)), hide_index=True, height=260, **WIDE)
        if len(ids) > 10:
            st.caption(f'{len(ids) - 10} more sheets are loaded and can be queried.')

for index, message in enumerate(st.session_state.messages):
    with st.chat_message(message['role']):
        st.markdown(message['content'])
        render_steps(message.get('steps'), f'm{index}')

if workbook and not st.session_state.messages and not st.session_state.pending_question:
    st.markdown('**Try asking**')
    options = suggestions(workbook)
    columns = st.columns(2)
    for index, question in enumerate(options):
        columns[index % 2].button(question, on_click=ask, args=(question,), key=f'suggestion_{index}', **WIDE)

typed = st.chat_input('Ask a question about your data…' if workbook else 'Upload a file to start', disabled=not workbook)
question = typed or st.session_state.pending_question
st.session_state.pending_question = None
if question and workbook:
    history = [{'role': m['role'], 'content': m['content']} for m in st.session_state.messages]
    with st.chat_message('user'):
        st.markdown(question)
    with st.chat_message('assistant'):
        started = time.monotonic()
        status = st.status('Thinking…', expanded=True)
        try:
            if not api_key:
                raise ValueError('Enter your GroqCloud API key in the sidebar.')
            chat = bounded_chat(history + [{'role': 'user', 'content': question}])
            model = build_model(model_name.strip(), api_key)
            answer = asyncio.run(answer_question(model, workbook, chat, status.write, input_budget, sheets=scope))
            queries = sum(step.frame is not None for step in answer.steps)
            status.update(label=f'Answered in {time.monotonic() - started:.1f}s · {queries} '
                                f'quer{"y" if queries == 1 else "ies"} over your data',
                          state='complete', expanded=False)
            st.markdown(answer.text)
            render_steps(answer.steps, f'm{len(st.session_state.messages) + 1}')
            st.session_state.messages += [{'role': 'user', 'content': question},
                                          {'role': 'assistant', 'content': answer.text, 'steps': answer.steps}]
        except Exception as exc:
            status.update(label='Stopped', state='error', expanded=False)
            st.error(friendly_error(exc))
