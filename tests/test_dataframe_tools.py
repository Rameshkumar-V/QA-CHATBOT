import io
import json
import unittest

import numpy as np
import pandas as pd

from services.data_tools import read_upload
from services.dataframe_tools import MAX_MODEL_ROWS, MAX_OUTPUT_CHARS, Toolkit
from services.table_store import MAX_RESULT_ROWS, QueryExecutionError
from services.workbook import Workbook, sheet_id


def sales(rows=10_000):
    rng = np.random.default_rng(7)
    return pd.DataFrame({'Order ID': range(1, rows + 1),
                         'Region': rng.choice(['East', 'West', 'North'], rows),
                         'Amount': rng.integers(10, 1000, rows).astype(float),
                         'Order Date': pd.date_range('2024-01-01', periods=rows, freq='h'),
                         'Notes': ['ok'] * (rows - 3) + [None, '  ', 'N/A']})


class ToolkitTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.frame = sales()
        cls.workbook = Workbook({'sales.xlsx :: Orders': cls.frame, 'regions.csv': pd.DataFrame(
            {'Region': ['East', 'West', 'North'], 'Manager': ['Asha', 'Ben', 'Chen']})})

    @classmethod
    def tearDownClass(cls):
        cls.workbook.close()

    def setUp(self):
        self.toolkit = Toolkit(self.workbook)
        self.tools = {tool.name: tool for tool in self.toolkit.tools()}

    def call(self, name, **kwargs):
        raw = self.tools[name].invoke(kwargs)
        self.assertLessEqual(len(raw), MAX_OUTPUT_CHARS)
        return json.loads(raw)

    def test_sheet_ids_are_short_and_sql_safe(self):
        self.assertEqual(list(self.workbook.frames), ['sales_orders', 'regions'])
        self.assertEqual(sheet_id('2024 Report.xlsx :: Order', set()), 't_2024_report_order')
        self.assertEqual(sheet_id('order.csv', set()), 'order_data')
        self.assertEqual(sheet_id('a.csv', {'a'}), 'a_2')

    def test_find_rows_counts_every_row_but_returns_few(self):
        result = self.call('find_rows', sheet='sales_orders', filters=[
            {'column': 'region', 'op': '==', 'value': 'EAST'}, {'column': 'Amount', 'op': '>', 'value': 900}],
            columns=['Order ID', 'Amount'], sort_by='Amount', descending=True, limit=500)
        expected = self.frame[(self.frame.Region == 'East') & (self.frame.Amount > 900)]
        self.assertEqual(result['matched'], len(expected))
        self.assertEqual(result['rows_scanned'], 10_000)
        self.assertEqual(len(result['rows']), MAX_MODEL_ROWS)
        self.assertEqual(result['rows'][0][1], expected.Amount.max())
        self.assertIn('user sees all', result['note'])
        self.assertEqual(len(self.toolkit.steps[-1].frame), len(expected))

    def test_aggregate_groups_all_rows(self):
        result = self.call('aggregate', sheet='sales_orders', group_by=['Region'],
                           metrics=[{'op': 'sum', 'column': 'Amount'}, {'op': 'count'}])
        totals = self.frame.groupby('Region').Amount.sum()
        self.assertEqual({row[0]: row[1] for row in result['rows']}, totals.astype(int).to_dict())
        self.assertEqual(sum(row[2] for row in result['rows']), 10_000)
        self.assertEqual(result['rows'][0][1], totals.max())
        self.assertEqual(self.toolkit.steps[-1].chart, 'bar')

    def test_monthly_trend_and_date_ranges(self):
        result = self.call('aggregate', sheet='sales_orders', group_by=['Order Date'], time_grain='month',
                           metrics=[{'op': 'sum', 'column': 'Amount'}])
        self.assertEqual(result['rows'][0][0], '2024-01')
        self.assertEqual(result['groups'], 14)
        february = self.call('aggregate', sheet='sales_orders', metrics=[{'op': 'count'}],
                             filters=[{'column': 'Order Date', 'op': 'between', 'values': ['2024-02-01', '2024-02-29']}])
        self.assertEqual(february['result']['count'], 29 * 24)

    def test_missing_values_and_text_operators(self):
        result = self.call('find_rows', sheet='sales_orders', filters=[{'column': 'Notes', 'op': 'is_missing'}],
                           columns=['Order ID'])
        self.assertEqual(result['matched'], 3)
        self.assertEqual(self.call('aggregate', sheet='sales_orders', metrics=[{'op': 'count', 'column': 'Notes'}])
                         ['result']['count(Notes)'], 9_997)
        result = self.call('find_rows', sheet='sales_orders', filters=[
            {'column': 'Region', 'op': 'in', 'values': ['east', 'west']}], match='all')
        self.assertEqual(result['matched'], int(self.frame.Region.isin(['East', 'West']).sum()))

    def test_describe_reports_types_ranges_and_categories(self):
        result = self.call('describe_sheet', sheet='sales_orders')
        columns = {c['name']: c for c in result['columns']}
        self.assertEqual(columns['Amount']['type'], 'number')
        self.assertEqual(columns['Order Date']['type'], 'date')
        self.assertEqual(set(columns['Region']['top_values']), {'East', 'West', 'North'})
        self.assertEqual(columns['Notes']['missing'], 3)
        self.assertEqual(len(result['sample_rows']['rows']), 3)

    def test_read_rows_is_capped(self):
        result = self.call('read_rows', sheet='sales_orders', start=9_999, count=1_000)
        self.assertEqual(result['shown'], 2)
        self.assertEqual(len(self.call('read_rows', sheet='sales_orders', count=1_000)['rows']), MAX_MODEL_ROWS)

    def test_unknown_column_suggests_alternatives_instead_of_failing(self):
        result = self.call('find_rows', sheet='sales_orders', filters=[{'column': 'Regon', 'op': '=', 'value': 'x'}])
        self.assertIn("'Region'", result['error'])
        self.assertTrue(self.toolkit.steps[-1].error)

    def test_sum_of_text_column_is_rejected(self):
        result = self.call('aggregate', sheet='sales_orders', metrics=[{'op': 'sum', 'column': 'Region'}])
        self.assertIn('no numeric values', result['error'])

    def test_sql_joins_sheets_with_readable_names(self):
        result = self.call('sql_query', sql='SELECT r.Manager, SUM(s.Amount) AS total FROM sales_orders s '
                                            'JOIN regions r ON s.Region = r.Region GROUP BY r.Manager ORDER BY total DESC')
        self.assertEqual(result['rows_returned'], 3)
        self.assertEqual(result['rows'][0][1], self.frame.groupby('Region').Amount.sum().max())
        self.assertTrue(self.toolkit.steps[-1].sql)

    def test_loose_arguments_from_small_models_still_work(self):
        result = self.call('aggregate', metrics='sum(Amount)', group_by='Region', descending='true', limit='2',
                           unexpected='ignored', sheet='Orders')
        self.assertEqual(result['shown'], 2)
        result = self.call('find_rows', filters=[{'column': 'notes', 'operator': 'is null'}], columns='Order ID',
                           sheet='sales_orders')
        self.assertEqual(result['matched'], 3)
        self.assertIn('empty in every returned column', self.call(
            'find_rows', sheet='sales_orders', filters=[{'column': 'Notes', 'op': 'is_missing'}], columns=['Notes'])['warning'])
        self.assertIn('unknown op', self.call('find_rows', sheet='sales_orders',
                                              filters=[{'column': 'Amount', 'op': 'approx', 'value': 1}])['error'])

    def test_scope_limits_sheets(self):
        toolkit = Toolkit(self.workbook, ['regions'])
        tools = {tool.name: tool for tool in toolkit.tools()}
        self.assertIn('error', json.loads(tools['sql_query'].invoke({'sql': 'SELECT COUNT(*) AS n FROM sales_orders'})))
        self.assertNotIn('sales_orders', self.workbook.overview(toolkit.ids))


class SqlSafetyTests(unittest.TestCase):
    def setUp(self):
        self.workbook = Workbook({'data.csv': pd.DataFrame({'Value': range(MAX_RESULT_ROWS + 2)})})

    def tearDown(self):
        self.workbook.close()

    def test_writes_and_external_access_are_rejected(self):
        for sql in ['DELETE FROM data', 'DROP TABLE data', 'UPDATE data SET Value=0',
                    "ATTACH DATABASE ':memory:' AS other", 'PRAGMA database_list',
                    "SELECT load_extension('x') FROM data", 'SELECT * FROM sqlite_master',
                    'SELECT * FROM data; DELETE FROM data', 'SELECT 42 AS invented']:
            with self.subTest(sql=sql), self.assertRaises(QueryExecutionError):
                self.workbook.store.query(sql)
        self.assertEqual(self.workbook.store.query('SELECT COUNT(*) AS n FROM data').frame.iloc[0]['n'], MAX_RESULT_ROWS + 2)

    def test_large_listings_are_marked_partial(self):
        result = self.workbook.store.query('SELECT Value FROM data')
        self.assertEqual(len(result.frame), MAX_RESULT_ROWS)
        self.assertTrue(result.truncated)

    def test_expensive_queries_time_out(self):
        with self.assertRaisesRegex(QueryExecutionError, 'timed out'):
            self.workbook.store.query('WITH RECURSIVE x(n) AS (SELECT Value FROM data UNION ALL SELECT n+1 FROM x) '
                                      'SELECT SUM(n) FROM x', timeout=0.01)


class LoaderTests(unittest.TestCase):
    def upload(self, text, name='data.csv'):
        file = io.BytesIO(text.encode('utf-8'))
        file.name = name
        return next(iter(read_upload(file).values()))

    def test_semicolon_csv_dates_and_formatted_numbers(self):
        frame = self.upload('Name;Joined;Fee\nA;31/01/2024;"₹1,200"\nB;05/02/2024;₹950\n')
        self.assertEqual(frame.columns.tolist(), ['Name', 'Joined', 'Fee'])
        self.assertEqual(frame['Joined'].dt.month.tolist(), [1, 2])
        self.assertEqual(frame['Fee'].tolist(), [1200, 950])

    def test_ambiguous_dates_and_ids_stay_text(self):
        frame = self.upload('When,Code\n01/02/2024,"0012,5"\n03/04/2024,"0099,1"\n')
        self.assertFalse(pd.api.types.is_datetime64_any_dtype(frame['When']))
        self.assertEqual(frame['Code'].iloc[0], '0012,5')

    def test_case_insensitive_duplicate_headers_are_unique(self):
        frame = self.upload('Name,name,NAME\n1,2,3\n')
        self.assertEqual(frame.columns.tolist(), ['Name', 'name (2)', 'NAME (3)'])

    def test_footer_notes_and_signatures_are_not_records(self):
        from openpyxl import Workbook as Book
        book = Book()
        book.active.append(['Register of students'])
        book.active.append(['S.No', 'Register Number', 'Name', 'Aadhar Number', 'Date of Birth', 'Signature'])
        for n in range(1, 30):
            book.active.append([n, f'R{n:03}', f'Student {n}', f'8235 9649 {8000 + n}', f'{n % 28 + 1:02}/05/2004'])
        book.active.append([None, 'Details are to be checked and certified by the Class in charge.'])
        book.active.append([None, None, None, 'Signature of the Class incharge'])
        upload = io.BytesIO()
        book.save(upload)
        upload.name = 'class.xlsx'
        frame = next(iter(read_upload(upload).values()))
        self.assertEqual(len(frame), 29)
        self.assertEqual(len(frame.attrs['notes']), 2)
        self.assertNotIn('Signature', frame.columns)  # empty once the footer is removed
        self.assertTrue(pd.api.types.is_datetime64_any_dtype(frame['Date of Birth']))
        self.assertEqual(frame['Aadhar Number'].iloc[0], '8235 9649 8001')  # identifier stays text
        self.assertTrue(pd.api.types.is_numeric_dtype(frame['S.No']))
        workbook = Workbook({'class.xlsx :: Sheet': frame})
        overview = workbook.overview()
        self.assertIn('29 data rows', overview)
        self.assertIn('not data rows', overview)
        workbook.close()

    def test_other_formats_are_rejected(self):
        with self.assertRaisesRegex(ValueError, 'Excel'):
            self.upload('{}', 'data.json')


if __name__ == '__main__':
    unittest.main()
