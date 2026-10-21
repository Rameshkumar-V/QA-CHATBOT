import io
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from services.data_tools import excel_previews, read_upload
from services.excel_headers import column_names, read_sheet
from services.workbook import Workbook as SheetBook
import pandas as pd


def workbook_upload(rows, merges=(), name='report.xlsx'):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    for cells in merges:
        book.active.merge_cells(cells)
    upload = io.BytesIO()
    book.save(upload)
    upload.name = name
    upload.size = len(upload.getvalue())
    upload.file_id = 'header-regression'
    return upload


class ExcelHeaderTests(unittest.TestCase):
    def test_merged_title_blank_rows_then_real_headers(self):
        labels = ['Record', 'Description'] + [f'Field {i}' for i in range(7)]
        rows = [['Arbitrary report title'], [], labels] + [[n, f'Entry {n}'] + [n] * 7 for n in range(1, 44)]
        upload = workbook_upload(rows, ['A1:I1'])
        tables = read_upload(upload)
        frame = tables['report.xlsx :: Sheet']
        self.assertEqual(frame.columns.tolist(), labels)
        self.assertEqual(frame.shape, (43, 9))
        self.assertEqual(frame.attrs['header_row'], 3)
        workbook = SheetBook(tables)
        self.assertEqual(workbook.frames['report_sheet'].columns.tolist(), labels)
        self.assertEqual(workbook.store.query('SELECT COUNT(*) AS n FROM report_sheet').frame.iloc[0]['n'], 43)
        workbook.close()

    def test_regular_text_only_table_does_not_lose_first_record(self):
        upload = workbook_upload([['Product', 'Location'], ['Pencil', 'North'], ['Eraser', 'South']])
        frame = next(iter(read_upload(upload).values()))
        self.assertEqual(frame.columns.tolist(), ['Product', 'Location'])
        self.assertEqual(frame.iloc[0]['Product'], 'Pencil')
        self.assertEqual(len(frame), 2)

    def test_blank_headers_keep_data_and_empty_columns_are_removed(self):
        upload = workbook_upload([['Code', None, None, 'Quantity'], [1, 'unlabelled', None, 10], [2, 'other', None, 20]])
        frame = next(iter(read_upload(upload).values()))
        self.assertEqual(frame.columns.tolist(), ['Code', 'Column B', 'Quantity'])
        self.assertEqual(frame.iloc[0]['Column B'], 'unlabelled')
        self.assertEqual(frame.attrs['empty_columns_removed'], 1)

    def test_duplicate_and_multiline_headers_are_clean_unique_and_stable(self):
        upload = workbook_upload([['  Label\n name ', 'Amount', 'Amount', 'Amount (2)'], ['a', 1, 2, 3]])
        frame = next(iter(read_upload(upload).values()))
        self.assertEqual(frame.columns.tolist(), ['Label name', 'Amount', 'Amount (3)', 'Amount (2)'])
        self.assertEqual(column_names(frame.columns), frame.columns.tolist())

    def test_manual_merged_multilevel_headers(self):
        upload = workbook_upload([['Code', 'Primary', None, 'Secondary', None],
                                  [None, 'Email', 'Phone', 'Email', 'Phone'],
                                  [1, 'a', 'b', 'c', 'd']], ['A1:A2', 'B1:C1', 'D1:E1'])
        frame = next(iter(read_upload(upload, excel_headers={'Sheet': {'header_row': 0, 'header_rows': 2}}).values()))
        self.assertEqual(frame.columns.tolist(), ['Code', 'Primary / Email', 'Primary / Phone', 'Secondary / Email', 'Secondary / Phone'])
        self.assertEqual(len(frame), 1)

    def test_numeric_headerless_data_retains_first_row(self):
        upload = workbook_upload([[1, 10], [2, 20], [3, 30]])
        frame = next(iter(read_upload(upload).values()))
        self.assertEqual(frame.columns.tolist(), ['Column A', 'Column B'])
        self.assertEqual(len(frame), 3)
        self.assertIsNone(frame.attrs['header_row'])
        self.assertEqual(frame.iloc[0]['Column A'], 1)

    def test_explicit_no_header_retains_text_records(self):
        upload = workbook_upload([['Alpha', 'North'], ['Beta', 'South']])
        frame = next(iter(read_upload(upload, excel_headers={'Sheet': {'header_row': None}}).values()))
        self.assertEqual(len(frame), 2)
        self.assertEqual(frame.iloc[0]['Column A'], 'Alpha')

    def test_each_sheet_has_independent_header_detection(self):
        book = Workbook()
        book.active.append(['Code', 'Value'])
        book.active.append([1, 10])
        other = book.create_sheet('Other')
        other.append(['Other report'])
        other.append([])
        other.append(['Key', 'Measure'])
        other.append([2, 20])
        upload = io.BytesIO()
        book.save(upload)
        upload.name = 'multiple.xlsx'
        tables = read_upload(upload)
        self.assertEqual(tables['multiple.xlsx :: Sheet'].columns.tolist(), ['Code', 'Value'])
        self.assertEqual(tables['multiple.xlsx :: Other'].columns.tolist(), ['Key', 'Measure'])

    def test_source_row_limit_does_not_silently_truncate(self):
        upload = workbook_upload([['Code', 'Value']] + [[i, i] for i in range(4)])
        with pd.ExcelFile(upload, engine='openpyxl') as book:
            with self.assertRaisesRegex(ValueError, 'maximum 3'):
                read_sheet(book, 'Sheet', max_rows=3)

    def test_preview_preserves_original_rows_for_manual_selection(self):
        upload = workbook_upload([['Title'], [], ['Code', 'Value'], [1, 2]])
        preview = excel_previews(upload)['Sheet']
        self.assertEqual(preview.iloc[0, 0], 'Title')
        self.assertEqual(preview.iloc[2, 0], 'Code')

    def test_ui_header_change_rebuilds_agent_schema_and_clears_stale_answers(self):
        upload = workbook_upload([['Report'], [], ['Code', 'Value'], [1, 2]], ['A1:B1'])
        with patch('streamlit.file_uploader', return_value=[upload]):
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py')).run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.session_state['workbook'].frames['report_sheet'].columns.tolist(), ['Code', 'Value'])
            app.session_state['messages'] = [{'role': 'user', 'content': 'Old schema question'},
                                             {'role': 'assistant', 'content': 'Old schema answer'}]
            next(w for w in app.selectbox if w.label == 'Header mode').select('Select header row').run()
            self.assertEqual(len(app.exception), 0)
            self.assertFalse(app.session_state['messages'])
            next(w for w in app.number_input if w.label == 'Header starts at Excel row').set_value(3).run()
            self.assertEqual(len(app.exception), 0)
            frame = app.session_state['workbook'].frames['report_sheet']
            self.assertEqual(frame.columns.tolist(), ['Code', 'Value'])
            self.assertEqual(frame.attrs['header_row'], 3)
