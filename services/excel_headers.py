"""Detect and normalize spreadsheet headers without knowing any field names."""
import re

import pandas as pd
from openpyxl.utils import get_column_letter

PREVIEW_ROWS = 30


def blank(value):
    return bool(pd.api.types.is_scalar(value) and pd.isna(value)) or isinstance(value, str) and not value.strip()


def label(value):
    return '' if blank(value) else re.sub(r'\s+', ' ', str(value)).strip()


def column_names(values):
    """Use real labels, make duplicates unique and label only genuinely blank headers."""
    base = [label(value) for value in values]
    base = [f'Column {get_column_letter(i + 1)}' if not name or re.fullmatch(r'Unnamed: \d+(?:_level_\d+)?', name)
            else name for i, name in enumerate(base)]
    reserved, seen, result = set(base), set(), []
    for name in base:
        candidate, suffix = name, 2
        while candidate in seen or (candidate != name and candidate in reserved):
            candidate = f'{name} ({suffix})'
            suffix += 1
        seen.add(candidate)
        result.append(candidate)
    return result


def _header_text(value):
    return isinstance(value, str) and bool(label(value)) and any(c.isalpha() for c in value)


def detect_header(preview):
    """Prefer the earliest dense label row over sparse titles and numeric data rows."""
    populated = [(index, list(row)) for index, row in enumerate(preview.itertuples(index=False, name=None))
                 if any(not blank(value) for value in row)]
    if not populated:
        return None
    active = [i for i in range(len(preview.columns)) if any(not blank(row[i]) for _, row in populated)]
    width = len(active)
    first_index, first_row = populated[0]
    first_values = [first_row[i] for i in active if not blank(first_row[i])]
    # A numeric first record is not silently discarded as a made-up header.
    if len(first_values) >= max(1, width * 0.6) and sum(_header_text(v) for v in first_values) / len(first_values) < 0.5:
        return None
    if len(first_values) >= max(2, width * 0.5) and all(_header_text(value) for value in first_values):
        return first_index
    scores = []
    for index, row in populated:
        values = [row[i] for i in active if not blank(row[i])]
        texts = sum(_header_text(value) for value in values)
        if texts < min(2, width) or texts / len(values) < 0.6:
            continue
        density = len(values) / width
        uniqueness = len({label(value).casefold() for value in values}) / len(values)
        long_cells = sum(len(label(value)) > 100 for value in values) / len(values)
        score = 3 * density + 2 * texts / len(values) + uniqueness - long_cells - 0.02 * index
        scores.append((score, index))
    if not scores:
        return None
    best_score, best_index = max(scores, key=lambda pair: (pair[0], -pair[1]))
    first_score = next((score for score, index in scores if index == first_index), None)
    if first_score is not None and best_score - first_score < 0.8:
        return first_index
    return best_index


def preview_sheet(book, sheet):
    return pd.read_excel(book, sheet_name=sheet, header=None, nrows=PREVIEW_ROWS, keep_default_na=False)


def read_sheet(book, sheet, max_rows, header_row='auto', header_rows=1):
    preview = preview_sheet(book, sheet)
    mode = 'automatic' if header_row == 'auto' else 'none' if header_row is None else 'manual'
    if header_row == 'auto':
        header_row = detect_header(preview)
        header_rows = 1
    if header_row is not None and (not isinstance(header_row, int) or header_row < 0 or header_row > 999):
        raise ValueError('Header row must be between Excel rows 1 and 1,000.')
    if header_rows not in (1, 2, 3):
        raise ValueError('Select one to three header rows.')
    if header_row is None:
        header = pd.DataFrame()
        skip = 0
    else:
        header = pd.read_excel(book, sheet_name=sheet, header=None, skiprows=header_row,
                               nrows=header_rows, keep_default_na=False)
        if header.empty:
            raise ValueError(f'{sheet}: the selected header row is beyond the sheet contents.')
        skip = header_row + header_rows
    frame = pd.read_excel(book, sheet_name=sheet, header=None, skiprows=skip,
                          nrows=max_rows + 1, keep_default_na=False, na_values=[''])
    if len(frame) > max_rows:
        raise ValueError(f'{sheet}: maximum {max_rows:,} source rows after headers')
    width = max(len(frame.columns), len(header.columns))
    frame = frame.reindex(columns=range(width))
    header = header.reindex(columns=range(width))
    if header_rows > 1 and not header.empty:
        for row in range(len(header) - 1):
            last = ''
            for col in header.columns:
                value = label(header.iat[row, col])
                if value:
                    last = value
                elif last:
                    header.iat[row, col] = last
    labels = []
    for col in range(width):
        parts = list(dict.fromkeys(label(v) for v in header[col] if label(v))) if not header.empty else []
        labels.append(' / '.join(parts))
    # A blank header with data must survive. Only completely empty columns are dropped.
    keep = [col for col in range(width) if labels[col] or any(not blank(v) for v in frame[col])]
    frame = frame[keep]
    names = column_names([labels[col] for col in keep])
    # Keep fallback labels tied to their actual Excel positions even after dropping empty columns.
    for index, source_col in enumerate(keep):
        if not labels[source_col]:
            names[index] = f'Column {get_column_letter(source_col + 1)}'
    frame.columns = column_names(names)
    nonempty = frame.apply(lambda column: column.map(lambda value: not blank(value))).any(axis=1)
    frame = frame.loc[nonempty].reset_index(drop=True)
    frame.attrs.update(header_row=None if header_row is None else header_row + 1,
                       header_rows=0 if header_row is None else header_rows, header_mode=mode,
                       generic_headers=[str(frame.columns[i]) for i, col in enumerate(keep) if not labels[col]],
                       empty_columns_removed=width - len(keep))
    return frame
